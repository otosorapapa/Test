# -*- coding: utf-8 -*-
"""
共通ユーティリティ。
- RAW_DIR, DATA_DIR などのディレクトリ定義
- 依存が無くても動く PDF 抽出・Word 出力フォールバック
- 簡易採点、JSON 入出力、XP 保存など
"""
import json, re, zipfile
from pathlib import Path
from typing import Any

# ---- ディレクトリ定数 ----
BASE_DIR = Path(".")
DATA_DIR = BASE_DIR / "data"
RAW_DIR = DATA_DIR / "raw"
EXTRACTED_DIR = DATA_DIR / "extracted"
ATTEMPTS_DIR = DATA_DIR / "attempts"
EXPORT_DIR = BASE_DIR / "export"
LOGS_DIR = BASE_DIR / "logs"
for d in [DATA_DIR, RAW_DIR, EXTRACTED_DIR, ATTEMPTS_DIR, EXPORT_DIR, LOGS_DIR]:
    d.mkdir(parents=True, exist_ok=True)

# ---- やさしい日本語変換 ----
EASY_MAP = {
    "課題": "困りごと",
    "方策": "やり方",
    "目的": "めあて",
    "KPI": "めあて",
    "実装": "やり方",
    "評価": "ふりかえり",
    "前提": "はじめの条件",
    "効果": "よい変化",
    "リスク": "きけん",
    "ステークホルダー": "関係者",
    "ガバナンス": "きまりごと",
    "ベンチマーク": "くらべるめやす",
}

def to_easy(text: str) -> str:
    if not text:
        return ""
    out = text
    for k, v in EASY_MAP.items():
        out = re.sub(k, v, out)
    return out

def simplify_text(text: str) -> str:
    return to_easy(text or "")

# ---- JSON ユーティリティ ----
def load_json(p: Path, default: Any):
    try:
        if p.exists():
            return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        pass
    return default

def save_json(p: Path, obj: Any):
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")

# ---- 採点（キーワード照合） ----
def grade_answer(answer: str, rubrics: dict, checklists: dict):
    if not answer:
        return {"total": 0, "detail": {}, "hits": {}, "lacks": []}
    total = 0
    detail = {}
    hits = {}
    low = answer.lower()
    for view, conf in rubrics.items():
        kws = conf.get("keywords", [])
        weight = conf.get("weight", 1)
        hit_count = 0
        hit_words = []
        for kw in kws:
            if kw.lower() in low:
                hit_count += 1
                hit_words.append(kw)
        score = min(hit_count, conf.get("max_hits", len(kws))) * weight
        total += score
        detail[view] = score
        hits[view] = hit_words
    lacks = [item for item in checklists.get("must_include", []) if item.lower() not in low]
    return {"total": total, "detail": detail, "hits": hits, "lacks": lacks}

# ---- 外部出力 ----
def export_markdown(body: str, meta: dict, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    title = meta.get("title", "export")
    fname = f"{title}.md"
    content = f"# {title}\n\n" + (body or "")
    path = out_dir / fname
    path.write_text(content, encoding="utf-8")
    return path

def export_docx(body: str, meta: dict, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    title = meta.get("title", "export")
    path = out_dir / f"{title}.docx"
    try:
        from docx import Document  # type: ignore
        doc = Document()
        if title:
            doc.add_heading(title, level=1)
        for line in (body or "").splitlines():
            doc.add_paragraph(line if line.strip() else "")
        doc.save(path)
        return path
    except Exception:
        # Fallback: 最小 DOCX を Zip で生成
        lines = (body or "").splitlines()
        def _doc_xml(lines):
            ps = []
            if title:
                ps.append(f"<w:p><w:r><w:t>{title}</w:t></w:r></w:p>")
            for ln in lines:
                t = ln.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                ps.append(f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>")
            inner = "".join(ps)
            return (
                "<?xml version='1.0' encoding='UTF-8' standalone='yes'?><w:document "
                "xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
                + inner
                + "</w:body></w:document>"
            )
        content_types = (
            "<?xml version='1.0' encoding='UTF-8'?><Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'>"
            "<Default Extension='rels' ContentType='application/vnd.openxmlformats-package.relationships+xml'/>"
            "<Default Extension='xml' ContentType='application/xml'/>"
            "<Override PartName='/word/document.xml' ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'/>"
            "</Types>"
        )
        rels = (
            "<?xml version='1.0' encoding='UTF-8'?><Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
            "<Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument' Target='word/document.xml'/>"
            "</Relationships>"
        )
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", content_types)
            zf.writestr("_rels/.rels", rels)
            zf.writestr("word/document.xml", _doc_xml(lines))
        return path

# ---- PDF テキスト抽出 ----
def extract_pdf_text(pdf_path: Path) -> str:
    try:
        from pdfminer.high_level import extract_text  # type: ignore
        return extract_text(str(pdf_path)) or ""
    except Exception:
        return f"[extract_pdf_text] フォールバック：PDF抽出ライブラリが未導入です。ファイル名: {pdf_path.name}"

# ---- 構成の分布 ----
def section_distribution(text: str) -> dict:
    if not text:
        return {"total": 0, "counts": {}, "ratio": {}}
    sections = {"序論": 0, "本論": 0, "結論": 0, "その他": 0}
    cur = "その他"
    for raw in text.splitlines():
        line = raw.strip()
        if re.search("序論", line):
            cur = "序論"
            continue
        if re.search("本論", line):
            cur = "本論"
            continue
        if re.search("結論", line):
            cur = "結論"
            continue
        sections[cur] += len(line)
    total = sum(sections.values()) or 1
    ratio = {k: round(v / total, 3) for k, v in sections.items()}
    return {"total": total, "counts": sections, "ratio": ratio}

# ---- XP の簡易保存 ----
def load_state() -> dict:
    p = BASE_DIR / "game_state.json"
    return load_json(p, {"xp": 0, "level": 1, "streak": 0, "badges": []})

def save_state(state: dict):
    p = BASE_DIR / "game_state.json"
    save_json(p, state)
